"""XLSX risk register export (plan checklist: XLSX risk register)."""

import io
from datetime import UTC, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.reports.data import ReviewData

HEADERS = [
    "ID",
    "Severity",
    "Category",
    "Clause",
    "Finding",
    "Exposure",
    "DBS Position",
    "Required Action",
    "Playbook Rule",
    "Confidence",
    "Evidence",
]

SEVERITY_FILL = {
    "critical": PatternFill("solid", start_color="F8D7DA"),
    "high": PatternFill("solid", start_color="FDE7CF"),
    "medium": PatternFill("solid", start_color="FFF3CD"),
    "low": PatternFill("solid", start_color="E8E8E8"),
    "informational": PatternFill("solid", start_color="EEEEEE"),
}

WIDTHS = [10, 12, 14, 34, 46, 14, 40, 40, 20, 10, 50]


def build_xlsx(data: ReviewData) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Risk Register"

    # Title block
    ws.append([f"DBS Contract Review — Risk Register — {data.document.filename}"])
    ws.append([f"Generated {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')}"])
    ws.append([])

    ws.append(HEADERS)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="1F3B63")

    for r in data.register:
        ws.append(
            [
                r.display_id,
                r.severity,
                r.category,
                r.clause,
                r.finding,
                r.exposure,
                r.dbs_position,
                r.required_action,
                r.rule or "",
                r.confidence,
                r.source_text,
            ]
        )
        row = ws[ws.max_row]
        fill = SEVERITY_FILL.get(r.severity)
        if fill:
            for cell in row:
                cell.fill = fill

    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{header_row + 1}"

    # Second sheet: exposures
    ws2 = wb.create_sheet("Financial Exposure")
    ws2.append(["Scenario", "Result", "Currency", "Formula", "Assumption", "Source"])
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="1F3B63")
    for e in data.exposures:
        ws2.append(
            [e.scenario, float(e.result), e.currency, e.formula, e.assumption, e.source_reference]
        )
    for i, w in enumerate([30, 16, 10, 50, 60, 24], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # Third sheet: cannot-calculate notes
    if data.cannot_calculate:
        ws3 = wb.create_sheet("Missing Data")
        ws3.append(["Metric", "Reason"])
        for cell in ws3[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", start_color="1F3B63")
        for n in data.cannot_calculate:
            ws3.append([n["metric"], n["reason"]])
        ws3.column_dimensions["A"].width = 40
        ws3.column_dimensions["B"].width = 50

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

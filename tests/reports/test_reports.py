import io
import uuid
import zipfile

import pytest
from docx import Document as DocxReader
from openpyxl import load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from sqlalchemy import text
from app.storage.database import engine


@pytest.fixture(scope="module", autouse=True)
async def _clean_db_module():
    """Clear database at module start for isolation."""
    async with engine.begin() as conn:
        await conn.execute(
            text("TRUNCATE TABLE audit_logs, users, findings, obligations, economic_exposures, contract_facts, extracted_tables, clauses, document_pages, document_sheets, documents RESTART IDENTITY CASCADE")
        )

PDF_TEXT = """\
1. COMMERCIAL
CLIENT: Omega Energy Limited
Contract Value: USD 20,000,000
Payment terms: within 75 days of invoice
Day Rate: USD 25,000

2. TERM
This Agreement shall commence on 1 January 2026 and continue for 36 months.

4. LIQUATED DAMAGES
Liquidated damages of USD 10,000 per day shall apply for each day of delay.
"""


def make_pdf() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    y = 740
    for line in PDF_TEXT.splitlines():
        c.drawString(60, y, line)
        y -= 16
    c.showPage()
    c.save()
    return buf.getvalue()


@pytest.fixture
async def analyzed_doc(client) -> str:
    resp = await client.post(
        "/api/v1/documents/upload",
        files={"file": ("review.pdf", make_pdf(), "application/pdf")},
    )
    doc_id = resp.json()["id"]
    await client.post(f"/api/v1/documents/{doc_id}/extract")
    yield doc_id
    # Cleanup
    await client.delete(f"/api/v1/documents/{doc_id}")


class TestReportDownloads:
    async def test_docx_report(self, client, analyzed_doc):
        resp = await client.get(f"/api/v1/documents/{analyzed_doc}/report.docx")
        assert resp.status_code == 200
        assert (
            resp.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        assert 'filename="' in resp.headers["content-disposition"]
        assert resp.content[:2] == b"PK"

        doc = DocxReader(io.BytesIO(resp.content))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert "DBS Contract Review Report" in [p.text for p in doc.paragraphs] or any(
            "Contract Review" in h for h in headings
        )
        assert any("Executive Summary" in h for h in headings)
        assert any("Risk Register" in h for h in headings)
        assert any("Financial Exposure" in h for h in headings)

    async def test_pdf_report(self, client, analyzed_doc):
        resp = await client.get(f"/api/v1/documents/{analyzed_doc}/report.pdf")
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/pdf"
        assert resp.content[:5] == b"%PDF-"
        assert len(resp.content) > 1000

    async def test_xlsx_register(self, client, analyzed_doc):
        resp = await client.get(f"/api/v1/documents/{analyzed_doc}/report.xlsx")
        assert resp.status_code == 200
        assert (
            resp.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        wb = load_workbook(io.BytesIO(resp.content))
        assert "Risk Register" in wb.sheetnames
        ws = wb["Risk Register"]
        headers = [c.value for c in ws[4]]
        assert headers[:4] == ["ID", "Severity", "Category", "Clause"]
        assert ws.max_row > 4
        assert "Financial Exposure" in wb.sheetnames

    async def test_invalid_format(self, client, analyzed_doc):
        resp = await client.get(f"/api/v1/documents/{analyzed_doc}/report.csv")
        assert resp.status_code == 400

    async def test_unknown_document(self, client):
        missing = uuid.uuid4()
        resp = await client.get(f"/api/v1/documents/{missing}/report.docx")
        assert resp.status_code == 404

    async def test_report_is_valid_zip(self, client, analyzed_doc):
        resp = await client.get(f"/api/v1/documents/{analyzed_doc}/report.docx")
        assert zipfile.is_zipfile(io.BytesIO(resp.content))

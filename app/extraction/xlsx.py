import logging
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)


@dataclass
class SheetResult:
    position: int
    sheet_name: str
    is_hidden: bool
    row_count: int
    col_count: int
    merged_ranges: list[str]
    data: list[dict[str, Any]]


@dataclass
class XlsxExtractionResult:
    sheets: list[SheetResult]
    warnings: list[str] = field(default_factory=list)

    @property
    def sheet_count(self) -> int:
        return len(self.sheets)


def _serialize_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float, str, bool)) or value is None:
        return value
    return str(value)


class XlsxExtractor:
    """Extract worksheets, cells, formulas, comments and hidden-sheet flags from XLSX."""

    def extract(self, xlsx_path: Path) -> XlsxExtractionResult:
        import openpyxl

        warnings: list[str] = []
        sheets: list[SheetResult] = []

        wb_formulas = openpyxl.load_workbook(xlsx_path, data_only=False, read_only=False)
        wb_values = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=False)

        hidden: list[str] = []
        for position, ws in enumerate(wb_formulas.worksheets, start=1):
            ws_values = wb_values[ws.title]
            is_hidden = ws.sheet_state != "visible"
            if is_hidden:
                hidden.append(f"{ws.title} ({ws.sheet_state})")

            merged_ranges = [str(r) for r in ws.merged_cells.ranges]
            data: list[dict[str, Any]] = []
            max_row = 0
            max_col = 0

            for row in ws.iter_rows():
                row_cells: list[dict[str, Any]] = []
                for cell in row:
                    formula_value = cell.value
                    comment = cell.comment.text if cell.comment else None
                    if formula_value is None and comment is None:
                        continue

                    is_formula = isinstance(formula_value, str) and formula_value.startswith("=")
                    cached = ws_values[cell.coordinate].value
                    value = _serialize_value(cached if is_formula else formula_value)

                    row_cells.append(
                        {
                            "ref": cell.coordinate,
                            "row": cell.row,
                            "col": cell.column,
                            "value": value,
                            "formula": formula_value if is_formula else None,
                            "comment": comment,
                        }
                    )
                    max_row = max(max_row, cell.row or 0)
                    max_col = max(max_col, cell.column or 0)

                if row_cells:
                    data.append(
                        {
                            "row": row[0].row,
                            "range": f"{row[0].coordinate}:{row[-1].coordinate}",
                            "cells": row_cells,
                        }
                    )

            sheets.append(
                SheetResult(
                    position=position,
                    sheet_name=ws.title,
                    is_hidden=is_hidden,
                    row_count=max_row,
                    col_count=max_col,
                    merged_ranges=merged_ranges,
                    data=data,
                )
            )

        if hidden:
            warnings.append("Hidden sheet(s) detected and flagged for review: " + ", ".join(hidden))
        wb_formulas.close()
        wb_values.close()
        return XlsxExtractionResult(sheets=sheets, warnings=warnings)

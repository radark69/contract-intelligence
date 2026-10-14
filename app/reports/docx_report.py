"""DOCX contract review report (plan section 5).

Professional internal-circulation report: cover, executive summary,
contract snapshot, risk register, critical findings, financial exposure,
obligation matrix and missing information.
"""

import io
from datetime import UTC, datetime

from docx import Document as Docx
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.table import Table

from app.reports.data import ReviewData

SEVERITY_COLOR = {
    "critical": RGBColor(0xB0, 0x00, 0x20),
    "high": RGBColor(0xD9, 0x6C, 0x00),
    "medium": RGBColor(0x8A, 0x6D, 0x00),
    "low": RGBColor(0x33, 0x33, 0x33),
    "informational": RGBColor(0x55, 0x55, 0x55),
}


def _fact_display(data: ReviewData, key: str, field: str | None = None) -> str:
    f = data.facts.get(key)
    if f is None or f.value is None:
        return "—"
    v = f.value
    if isinstance(v, dict):
        if field:
            v = v.get(field, "—")
        else:
            parts = [str(x) for x in v.values()]
            return " ".join(parts)
    return str(v)


def _add_heading(doc: DocxDocument, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _add_table(doc: DocxDocument, headers: list[str], rows: list[list[str]]) -> Table:
    table: Table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    return table


def build_docx(data: ReviewData) -> bytes:
    doc = Docx()
    d = data.document

    # --- Cover ---------------------------------------------------------------
    title = doc.add_heading("DBS Contract Review Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(d.filename)
    run.bold = True
    run.font.size = Pt(14)

    meta = [
        ("Client / Counterparty", d.client or "—"),
        ("Contractor", d.contractor or "—"),
        ("Contract type", d.contract_type or "—"),
        ("Review date", datetime.now(UTC).strftime("%Y-%m-%d")),
        ("Contract version", d.version or "—"),
        ("Reviewer", "DBS Contract Review System"),
        ("Classification", "Confidential — Internal DBS Circulation"),
    ]
    _add_table(doc, ["Field", "Value"], [[k, v] for k, v in meta])

    # --- 1. Executive summary -------------------------------------------------
    _add_heading(doc, "1. Executive Summary", 1)
    critical = data.severity_counts.get("critical", 0)
    high = data.severity_counts.get("high", 0)
    summary = [
        ("Contract value", _fact_display(data, "contract_value")),
        ("Duration", _fact_display(data, "duration")),
        ("Critical risks", str(critical)),
        ("High risks", str(high)),
        ("Findings total", str(len(data.findings_sorted))),
        ("Management approval required", ", ".join(data.approval_rules) or "None"),
        (
            "Missing information items",
            str(sum(1 for r in data.register if r.rule == "MISSING-INFO")),
        ),
    ]
    _add_table(doc, ["Metric", "Value"], [[k, v] for k, v in summary])

    key_issues = [r for r in data.register if r.severity in ("critical", "high")][:8]
    if key_issues:
        doc.add_paragraph("Key negotiation issues:")
        for r in key_issues:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{r.display_id} — ").bold = True
            p.add_run(f"{r.finding} ({r.required_action})")

    # --- 2. Contract snapshot -------------------------------------------------
    _add_heading(doc, "2. Contract Snapshot", 1)
    pt_days = data.facts.get("payment_terms_days")
    payment_display = (
        f"{pt_days.value} days" if pt_days is not None and pt_days.value is not None else "—"
    )
    avail = data.facts.get("availability_target")
    avail_display = f"{avail.value}%" if avail is not None and avail.value is not None else "—"
    snapshot = [
        ("Client", d.client or "—"),
        ("Contractor", d.contractor or "—"),
        ("Contract value", _fact_display(data, "contract_value")),
        ("Currency", d.currency or "—"),
        ("Duration", _fact_display(data, "duration")),
        ("Effective date", (d.start_date.strftime("%Y-%m-%d") if d.start_date else "—")),
        ("Expiry date", (d.end_date.strftime("%Y-%m-%d") if d.end_date else "—")),
        ("Payment terms", payment_display),
        ("Availability requirement", avail_display),
        ("Liability cap", _fact_display(data, "liability_cap")),
        ("LD rate", _fact_display(data, "ld_rate")),
        ("Day rate", _fact_display(data, "day_rate")),
        ("Clauses analysed", str(data.clause_count)),
    ]
    _add_table(doc, ["Field", "Value"], [[k, v] for k, v in snapshot])

    # --- 3. Risk register -----------------------------------------------------
    _add_heading(doc, "3. Risk Register", 1)
    if data.register:
        reg_table = _add_table(
            doc,
            [
                "ID",
                "Sev",
                "Category",
                "Clause",
                "Finding",
                "Exposure",
                "DBS Position",
                "Required Action",
            ],
            [
                [
                    r.display_id,
                    r.severity,
                    r.category,
                    r.clause[:60],
                    r.finding[:80],
                    f"{r.exposure:,.0f}" if r.exposure is not None else "—",
                    r.dbs_position[:80],
                    r.required_action[:80],
                ]
                for r in data.register
            ],
        )
        for row in reg_table.rows[1:]:
            cell = row.cells[1]
            color = SEVERITY_COLOR.get(cell.text.strip().lower())
            if color is None:
                continue
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.color.rgb = color
                    run.bold = True
    else:
        doc.add_paragraph("No findings recorded.")

    # --- 4. Critical & high findings in detail --------------------------------
    _add_heading(doc, "4. Critical Findings", 1)
    detailed = [r for r in data.register if r.severity in ("critical", "high")]
    if not detailed:
        doc.add_paragraph("No critical or high findings.")
    for r in detailed:
        _add_heading(doc, f"{r.display_id} — {r.finding}", level=2)
        rows = [
            ("Severity", r.severity),
            ("Category", r.category),
            ("Source", r.clause),
            ("Risk", r.dbs_position),
            ("Impact", r.impact or "—"),
            ("Action", r.required_action),
            ("Playbook rule", r.rule or "—"),
            ("Evidence", (r.source_text or "—")[:400]),
        ]
        _add_table(doc, ["Field", "Detail"], [[k, v] for k, v in rows])

    # --- 5. Financial exposure -------------------------------------------------
    _add_heading(doc, "5. Financial Exposure", 1)
    if data.exposures:
        _add_table(
            doc,
            ["Scenario", "Result", "Currency", "Formula", "Assumption"],
            [
                [e.scenario, f"{float(e.result):,.2f}", e.currency, e.formula, e.assumption[:120]]
                for e in data.exposures
            ],
        )
    if data.cannot_calculate:
        doc.add_paragraph("Cannot calculate reliably:")
        for n in data.cannot_calculate:
            doc.add_paragraph(
                f"Cannot calculate reliably — {n['metric']}: {n['reason']}.",
                style="List Bullet",
            )

    # --- 6. Obligation matrix ---------------------------------------------------
    _add_heading(doc, "6. Obligation Matrix", 1)
    if data.obligations:
        _add_table(
            doc,
            ["Party", "Obligation", "Trigger", "Frequency", "Source"],
            [
                [
                    o.party,
                    o.obligation[:120],
                    o.trigger or "—",
                    o.frequency or "—",
                    o.source_reference,
                ]
                for o in data.obligations
            ],
        )
    else:
        doc.add_paragraph("No obligations extracted.")

    # --- 7. Missing information -------------------------------------------------
    missing = [r for r in data.register if r.rule == "MISSING-INFO"]
    if missing:
        _add_heading(doc, "7. Missing Information / Clarifications", 1)
        for r in missing:
            doc.add_paragraph(f"{r.finding} — {r.required_action}", style="List Bullet")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
